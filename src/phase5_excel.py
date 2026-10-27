"""Phase 5 - build reports/ISEM_market_summary.xlsx for a trading audience (openpyxl).

Sheets: Summary (KPIs, formulas pointing at the detail sheets), Monthly (price statistics),
Heatmap (hour x month average price, colour scale), Battery (monthly revenue by strategy,
totals and capture by formula, native chart), Forecast (accuracy by model, native chart),
Notes (sources and caveats). All values come from the DuckDB database or the phase metrics JSON.
"""
import json

import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import DB_PATH, REPORTS

NAVY, BLUE, LIGHT, GREY = "0D366B", "2A78D6", "EAF2FC", "F3F3F1"
HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor=NAVY)
TOTAL_FILL = PatternFill("solid", fgColor=LIGHT)
THIN = Side(style="thin", color="D0D0CC")
BOX = Border(bottom=THIN)
EUR = '#,##0.0" €"'
EUR0 = '#,##0" €"'


def monthly_stats(con):
    d = con.execute("""
        SELECT h.trading_date, h.ts_local, h.price_eur_mwh AS p, h.wind_share, h.wind_all_mw, h.load_roi_mw,
               g.price_eur_mwh AS gb
        FROM hourly_panel h LEFT JOIN price_gb g USING (ts_utc)
        WHERE h.in_study ORDER BY h.ts_utc""").df()
    d["month"] = pd.to_datetime(d.trading_date).dt.to_period("M")
    day = d.dropna(subset=["p"]).groupby("trading_date").p.agg(lambda x: x.max() - x.min())
    day.index = pd.to_datetime(day.index).to_period("M")
    w = d.dropna(subset=["p", "wind_all_mw"])
    cap = (w.p * w.wind_all_mw).groupby(w.month).sum() / w.wind_all_mw.groupby(w.month).sum()
    g = d.groupby("month")
    m = pd.DataFrame({
        "Mean price": g.p.mean(), "Min": g.p.min(), "Max": g.p.max(), "Std dev": g.p.std(),
        "Hours <= 5 EUR": g.p.apply(lambda s: int((s <= 5).sum())),
        "Negative hours": g.p.apply(lambda s: int((s < 0).sum())),
        "Median daily spread": day.groupby(level=0).median(),
        "Wind capture price": cap,
        "Mean wind penetration": g.wind_share.mean(),
        "Mean all-island wind (MW)": g.wind_all_mw.mean(),
        "GB mean price": g.gb.mean(),
    })
    m.index = m.index.strftime("%b %Y")
    heat = d.dropna(subset=["p"]).assign(hour=d.ts_local.dt.hour).pivot_table(
        index="hour", columns="month", values="p", aggfunc="mean")
    heat.columns = [c.strftime("%b %y") for c in heat.columns]
    return m, heat


def header(ws, row, values, col=1):
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=col + j, value=v)
        c.font, c.fill = HDR_FONT, HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, size=9, color="52514E")


def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def main():
    con = duckdb.connect(str(DB_PATH), read_only=True)
    m, heat = monthly_stats(con)
    con.close()
    p3 = json.loads((REPORTS / "phase3_metrics.json").read_text())
    p4 = json.loads((REPORTS / "phase4_metrics.json").read_text())
    wb = Workbook()

    # ------------------------------------------------ Monthly
    ws = wb.active
    ws.title = "Monthly"
    title(ws, "I-SEM day-ahead prices by trading month (EUR/MWh)",
          "Study period Oct 2024 - Sep 2026. Wind capture price = wind-weighted average price. "
          "Wind penetration = ROI wind / ROI load.")
    cols = ["Month"] + list(m.columns)
    header(ws, 4, cols)
    for i, (idx, r) in enumerate(m.iterrows(), start=5):
        ws.cell(row=i, column=1, value=idx)
        for j, c in enumerate(m.columns, start=2):
            cell = ws.cell(row=i, column=j, value=None if pd.isna(r[c]) else float(r[c]))
            cell.number_format = "0.0%" if "penetration" in c else ("0" if "hours" in c.lower() else "#,##0.0")
            cell.border = BOX
    first, last = 5, 4 + len(m)
    for k, (lab, fn) in enumerate([("Average", "AVERAGE"), ("Minimum", "MIN"), ("Maximum", "MAX")]):
        row = last + 1 + k
        ws.cell(row=row, column=1, value=lab).font = Font(bold=True)
        for j in range(2, len(cols) + 1):
            L = get_column_letter(j)
            f = "SUM" if ("hours" in cols[j - 1].lower() and fn == "AVERAGE") else fn
            cell = ws.cell(row=row, column=j, value=f"={f}({L}{first}:{L}{last})")
            cell.number_format = ws.cell(row=first, column=j).number_format
            cell.fill = TOTAL_FILL
            cell.font = Font(bold=True)
    ws.cell(row=last + 1, column=1, value="Average (hours: total)")
    ws.freeze_panes = "B5"
    widths(ws, [22] + [13] * (len(cols) - 1))
    ws.conditional_formatting.add(f"B{first}:B{last}", ColorScaleRule(
        start_type="min", start_color="F7FBFF", end_type="max", end_color="2A78D6"))
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.height, ch.width = "Monthly mean price vs GB", "EUR/MWh", 8, 22
    ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=last), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=len(cols), min_row=4, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    ws.add_chart(ch, f"A{last + 6}")

    # ------------------------------------------------ Heatmap
    hs = wb.create_sheet("Heatmap")
    title(hs, "Average DAM price by hour (Irish local time) and trading month, EUR/MWh",
          "Conditional colour scale: blue = cheap, red = expensive.")
    header(hs, 4, ["Hour"] + list(heat.columns))
    for i, (h, r) in enumerate(heat.iterrows(), start=5):
        hs.cell(row=i, column=1, value=f"{h:02d}:00").font = Font(bold=True)
        for j, v in enumerate(r.values, start=2):
            c = hs.cell(row=i, column=j, value=round(float(v), 1))
            c.number_format = "0"
            c.alignment = Alignment(horizontal="center")
    rng = f"B5:{get_column_letter(1 + heat.shape[1])}{4 + len(heat)}"
    hs.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="2A78D6",
                                                      mid_type="percentile", mid_value=50, mid_color="F7F7F5",
                                                      end_type="max", end_color="E34948"))
    hs.freeze_panes = "B5"
    widths(hs, [8] + [7.5] * heat.shape[1])

    # ------------------------------------------------ Battery
    bs = wb.create_sheet("Battery")
    title(bs, "Battery DAM arbitrage revenue, 1 MW / 2 MWh, 88% RTE, test year (EUR per MW)",
          "Perfect foresight = upper bound. Forecast = optimised on the LightGBM forecast, settled at actual prices. "
          "Rule = fixed seasonal hours from Year 1. DAM only - excludes system services, capacity, intraday, balancing.")
    strat = [("perfect", "Perfect foresight"), ("forecast_lightgbm", "Forecast (LightGBM)"),
             ("forecast_naive", "Forecast (naive)"), ("rule", "Simple rule")]
    header(bs, 4, ["Month"] + [s[1] for s in strat] + ["Forecast capture", "Value of forecasting"])
    mon = p4["monthly_revenue_eur"]
    months = sorted(mon)
    for i, mo in enumerate(months, start=5):
        bs.cell(row=i, column=1, value=pd.Period(mo).strftime("%b %Y"))
        for j, (key, _) in enumerate(strat, start=2):
            c = bs.cell(row=i, column=j, value=float(mon[mo][key]))
            c.number_format = EUR0
        bs.cell(row=i, column=6, value=f"=C{i}/B{i}").number_format = "0%"
        bs.cell(row=i, column=7, value=f"=C{i}-E{i}").number_format = EUR0
    tr = 5 + len(months)
    bs.cell(row=tr, column=1, value="Total (scored days)").font = Font(bold=True)
    for j in range(2, 6):
        L = get_column_letter(j)
        c = bs.cell(row=tr, column=j, value=f"=SUM({L}5:{L}{tr - 1})")
        c.number_format, c.fill, c.font = EUR0, TOTAL_FILL, Font(bold=True)
    for j, f in [(6, f"=C{tr}/B{tr}"), (7, f"=C{tr}-E{tr}")]:
        c = bs.cell(row=tr, column=j, value=f)
        c.number_format, c.fill, c.font = ("0%" if j == 6 else EUR0), TOTAL_FILL, Font(bold=True)
    days = p4["days_in_backtest"]
    bs.cell(row=tr + 1, column=1, value=f"Annualised (x 365 / {days} days)").font = Font(bold=True)
    for j in range(2, 6):
        L = get_column_letter(j)
        c = bs.cell(row=tr + 1, column=j, value=f"={L}{tr}*365/{days}")
        c.number_format, c.font = EUR0, Font(bold=True)
    widths(bs, [22, 18, 20, 18, 14, 16, 20])
    bc = BarChart()
    bc.type, bc.grouping, bc.overlap = "col", "clustered", -10
    bc.title, bc.y_axis.title, bc.height, bc.width = "Monthly revenue by strategy", "EUR per MW", 8, 22
    for j in (2, 3, 5):
        bc.add_data(Reference(bs, min_col=j, min_row=4, max_row=tr - 1), titles_from_data=True)
    bc.set_categories(Reference(bs, min_col=1, min_row=5, max_row=tr - 1))
    bs.add_chart(bc, f"A{tr + 4}")
    # Sensitivity table
    sr = tr + 22
    bs.cell(row=sr, column=1, value="Sensitivity: forecast strategy, EUR per MW per year").font = Font(bold=True, color=NAVY)
    header(bs, sr + 1, ["Duration", "80% RTE", "88% RTE", "92% RTE"])
    sens = pd.DataFrame(p4["sensitivity"])
    for i, h in enumerate((1, 2, 4), start=sr + 2):
        bs.cell(row=i, column=1, value=f"{h} h ({h} MWh)")
        for j, r in enumerate((0.80, 0.88, 0.92), start=2):
            v = sens[(sens.duration_h == h) & (sens.rte == r) & (sens.strategy == "forecast_lightgbm")].annual_eur_per_mw
            bs.cell(row=i, column=j, value=float(v.iloc[0])).number_format = EUR0
    bs.conditional_formatting.add(f"B{sr + 2}:D{sr + 4}", ColorScaleRule(
        start_type="min", start_color="F7FBFF", end_type="max", end_color="2A78D6"))

    # ------------------------------------------------ Forecast
    fs = wb.create_sheet("Forecast")
    title(fs, "Day-ahead price forecast accuracy, test year (EUR/MWh)",
          "Information set = what is known at the D-1 11:00 gate closure. 'Upper bound' uses actual wind and is not a forecast.")
    header(fs, 4, ["Model", "MAE", "RMSE", "Bias (mean error)", "MAE vs naive"])
    names = [("naive_d7", "Naive: same hour last week"), ("naive_d1", "Naive: same hour yesterday"),
             ("linear", "Linear regression"), ("lightgbm", "LightGBM"), ("lightgbm_icon", "LightGBM (ICON weather)"),
             ("lightgbm_perfect_wind", "Upper bound: + actual wind")]
    for i, (k, lab) in enumerate(names, start=5):
        fs.cell(row=i, column=1, value=lab)
        fs.cell(row=i, column=2, value=p3["overall"][k]["MAE"]).number_format = "0.0"
        fs.cell(row=i, column=3, value=p3["overall"][k]["RMSE"]).number_format = "0.0"
        fs.cell(row=i, column=4, value=p3["mean_error"][k]).number_format = "+0.0;-0.0"
        fs.cell(row=i, column=5, value=f"=1-B{i}/$B$6").number_format = "0%"
    widths(fs, [32, 10, 10, 16, 14])
    fc = BarChart()
    fc.type, fc.title, fc.height, fc.width = "bar", "Mean absolute error by model", 7, 16
    fc.add_data(Reference(fs, min_col=2, min_row=4, max_row=10), titles_from_data=True)
    fc.set_categories(Reference(fs, min_col=1, min_row=5, max_row=10))
    fc.legend = None
    fs.add_chart(fc, "G4")
    hr = 13
    fs.cell(row=hr, column=1, value="MAE by delivery hour").font = Font(bold=True, color=NAVY)
    header(fs, hr + 1, ["Hour", "Naive (yesterday)", "Linear", "LightGBM"])
    for i, h in enumerate(range(24), start=hr + 2):
        fs.cell(row=i, column=1, value=f"{h:02d}:00")
        for j, k in enumerate(["naive_d1", "linear", "lightgbm"], start=2):
            fs.cell(row=i, column=j, value=p3["by_hour_mae"][k][str(h)]).number_format = "0.0"

    # ------------------------------------------------ Summary (first sheet)
    ss = wb.create_sheet("Summary", 0)
    title(ss, "I-SEM wind, price and battery arbitrage - summary",
          "Trading days 1 Oct 2024 - 30 Sep 2026; test year = Oct 2025 - Sep 2026. Source data: Ember/SEMOpx, "
          "Energy-Charts/ENTSO-E, Open-Meteo. See Notes.")
    header(ss, 4, ["KPI", "Value", "Unit", "Comment"])
    n = len(m)
    kpis = [
        ("Mean DAM price (study period)", f"=Monthly!B{5 + n}", "EUR/MWh", "Average of monthly means"),
        ("Highest monthly mean", f"=Monthly!B{7 + n}", "EUR/MWh", "Sep 2026 - the level tracks GB/gas"),
        ("Hours at or below 5 EUR/MWh", f"=Monthly!F{5 + n}", "hours", "1.5% of hours; mostly night and midday"),
        ("Median daily spread (max - min)", json.loads((REPORTS / 'phase2_metrics.json').read_text())["q4"]["study_median"],
         "EUR/MWh", "The battery revenue pool"),
        ("Merit-order effect", -json.loads((REPORTS / 'phase2_metrics.json').read_text())["q2"]["month_fe"]["slope_per_10pp"],
         "EUR/MWh per 10 pp", "Price fall per +10 points of wind penetration, within month"),
        ("Forecast MAE - LightGBM", "=Forecast!B8", "EUR/MWh", "Test year"),
        ("Forecast MAE - naive (yesterday)", "=Forecast!B6", "EUR/MWh", "Baseline"),
        ("Forecast improvement vs naive", "=Forecast!E8", "%", "Lower MAE is better"),
        ("Battery revenue - perfect foresight", f"=Battery!B{tr + 1}", "EUR/MW/yr", "Upper bound"),
        ("Battery revenue - forecast-based", f"=Battery!C{tr + 1}", "EUR/MW/yr", "Optimised on LightGBM"),
        ("Battery revenue - simple rule", f"=Battery!E{tr + 1}", "EUR/MW/yr", "Fixed seasonal hours"),
        ("Share of perfect foresight captured", f"=Battery!C{tr + 1}/Battery!B{tr + 1}", "%", "Forecast strategy"),
        ("Value of forecasting", f"=Battery!C{tr + 1}-Battery!E{tr + 1}", "EUR/MW/yr", "Forecast minus rule"),
    ]
    for i, (k, v, u, c) in enumerate(kpis, start=5):
        ss.cell(row=i, column=1, value=k).font = Font(bold=True)
        cell = ss.cell(row=i, column=2, value=v)
        cell.number_format = "0%" if u == "%" else ("#,##0" if u in ("EUR/MW/yr", "hours") else "0.0")
        cell.font = Font(bold=True, color=BLUE, size=12)
        ss.cell(row=i, column=3, value=u)
        ss.cell(row=i, column=4, value=c).font = Font(color="52514E")
        for j in range(1, 5):
            ss.cell(row=i, column=j).border = BOX
    widths(ss, [38, 16, 18, 62])
    ss.sheet_view.showGridLines = False

    # ------------------------------------------------ Notes
    ns = wb.create_sheet("Notes")
    title(ns, "Sources, definitions and caveats")
    notes = [
        "Prices: Ember European Wholesale Electricity Price Data (CC-BY-4.0, from SEMOpx), corrected -24 h; "
        "flat-price days replaced with official SEMOpx files or excluded (see docs/DECISIONS.md D9).",
        "Wind and load: Energy-Charts (Fraunhofer ISE), ENTSO-E data; ROI = EirGrid area, NI = SONI area. "
        "Validated against EirGrid Smart Grid Dashboard (corr 0.998).",
        "Weather forecasts: Open-Meteo Previous Runs API, ECMWF IFS 0.25, 48 h-ahead vintage, six locations (CC-BY-4.0).",
        "Wind penetration uses ROI wind / ROI load because NI load is unpublished for 55% of the period (D8).",
        "Trading day = 23:00-23:00 Irish local time. Hour columns are Irish local time.",
        "Battery figures are DAM-only arbitrage: no system services, capacity, intraday or balancing revenue; "
        "no degradation, network charges or imbalance costs; price-taker assumption.",
        "Every number is produced by code in src/ - see README for the one-command rebuild.",
    ]
    for i, t in enumerate(notes, start=3):
        c = ns.cell(row=i, column=1, value=f"- {t}")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ns.row_dimensions[i].height = 32
    ns.column_dimensions["A"].width = 130

    out = REPORTS / "ISEM_market_summary.xlsx"
    wb.save(out)
    print(f"saved {out}")


if __name__ == "__main__":
    main()
